import json
from io import BytesIO
from xml.sax.saxutils import escape

import pandas as pd
import requests
import streamlit as st
from anthropic import Anthropic
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable, ListFlowable, ListItem)

MODEL = "claude-haiku-4-5-20251001"
st.set_page_config(page_title="Resume Tailor", page_icon="📄", layout="centered")
client = Anthropic(api_key=st.secrets["ANTHROPIC_API_KEY"])

EMPTY_PROFILE = {
    "contact": {"name": "", "email": "", "phone": "", "location": "", "linkedin": "", "github": "", "website": ""},
    "summaries": {"default": ""},
    "experience": [], "projects": [], "education": [],
    "skills": {"all": []}, "certifications": [], "affiliations": [], "leadership": [],
}


# ================= Claude helpers =================

def ask_claude(prompt, max_tokens, images=None):
    content = prompt
    if images:
        content = [{"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": img}}
                   for img in images] + [{"type": "text", "text": prompt}]
    msg = client.messages.create(model=MODEL, max_tokens=max_tokens,
                                 messages=[{"role": "user", "content": content}])
    text = msg.content[0].text.strip().replace("```json", "").replace("```", "").strip()
    return json.loads(text)


def normalize(profile):
    """Make sure every field exists and has the right shape."""
    p = json.loads(json.dumps(EMPTY_PROFILE))
    for k in p:
        if profile.get(k):
            p[k] = profile[k]
    p["contact"] = {**EMPTY_PROFILE["contact"], **(profile.get("contact") or {})}
    for group in ["experience", "projects"]:
        for item in p[group]:
            item["bullets"] = [b if isinstance(b, dict) else {"text": str(b)} for b in item.get("bullets", [])]
    p["leadership"] = [l if isinstance(l, dict) else {"text": str(l)} for l in p["leadership"]]
    p["affiliations"] = [a if isinstance(a, dict) else {"name": str(a)} for a in p["affiliations"]]
    if not isinstance(p["skills"], dict):
        p["skills"] = {"all": list(p["skills"])}
    if not p["summaries"]:
        p["summaries"] = {"default": ""}
    return p


def parse_resume(resume_text="", images=None):
    if images:
        source = ("The images above are screenshots of this person's resume or LinkedIn profile. "
                  "Read all of them together. Ignore app buttons, menus, ads, and other people's names.")
    else:
        source = f"<resume>\n{resume_text}\n</resume>"
    prompt = f"""Convert this resume into JSON. Copy every bullet WORD FOR WORD.
Do not rewrite, improve, or invent anything. If a field is missing, use "" or [].

{source}

Respond with only JSON in this exact shape:
{{"contact": {{"name": "", "email": "", "phone": "", "location": "", "linkedin": "", "github": "", "website": ""}},
 "summaries": {{"default": "summary text from the resume, or empty"}},
 "experience": [{{"title": "", "company": "", "dates": "", "location": "", "bullets": [{{"text": ""}}]}}],
 "projects": [{{"name": "", "role": "", "dates": "", "bullets": [{{"text": ""}}]}}],
 "education": [{{"degree": "", "school": "", "graduation": "", "location": ""}}],
 "skills": {{"all": ["skill"]}},
 "certifications": [{{"name": "", "issuer": "", "date": ""}}],
 "affiliations": [{{"name": ""}}],
 "leadership": [{{"text": ""}}]}}
"""
    return normalize(ask_claude(prompt, 4000, images))


def build_catalog(profile):
    catalog, lines = {}, []
    for i, job in enumerate(profile["experience"]):
        lines.append(f"\n[JOB] {job['title']} at {job['company']} ({job['dates']})")
        for j, b in enumerate(job["bullets"]):
            catalog[f"exp{i}_{j}"] = b["text"]
            lines.append(f"exp{i}_{j}: {b['text']}")
    for i, proj in enumerate(profile["projects"]):
        lines.append(f"\n[PROJECT] {proj['name']} ({proj['dates']})")
        for j, b in enumerate(proj["bullets"]):
            catalog[f"proj{i}_{j}"] = b["text"]
            lines.append(f"proj{i}_{j}: {b['text']}")
    return catalog, "\n".join(lines)


def all_skills(profile):
    return [s for g in profile["skills"].values() for s in g]


def tailor(profile, catalog_text, job_posting):
    certs = "\n".join(f"cert{i}: {c['name']} ({c['issuer']})" for i, c in enumerate(profile["certifications"]))
    affs = "\n".join(f"aff{i}: {a['name']}" for i, a in enumerate(profile["affiliations"]))
    leads = "\n".join(f"lead{i}: {l['text']}" for i, l in enumerate(profile["leadership"]))
    prompt = f"""Pick the best resume content for this job posting.

<job_posting>
{job_posting}
</job_posting>

<summaries>
{json.dumps(profile["summaries"], indent=2)}
</summaries>

<bullets>
{catalog_text}
</bullets>

<certifications>
{certs}
</certifications>

<affiliations>
{affs}
</affiliations>

<leadership>
{leads}
</leadership>

Rules:
- Pick 12 to 15 bullet IDs total, or all of them if there are fewer.
- Give each chosen job at least 2 bullets when possible.
- Pick the best summary key.
- Pick up to 10 skills from this list only: {json.dumps(all_skills(profile))}
- Pick up to 3 certification IDs, up to 5 affiliation IDs, and up to 2 leadership IDs that fit this job.

Respond with only JSON in this exact shape, and nothing else:
{{"summary_key": "default", "bullet_ids": ["exp0_1"], "cert_ids": ["cert0"], "affiliation_ids": ["aff0"], "leadership_ids": ["lead0"], "skills_to_highlight": ["Python"], "why": "one sentence on your choices"}}
"""
    return ask_claude(prompt, 1000)


# ================= File and link reading =================

def read_upload(file):
    name = file.name.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(file).pages)
    if name.endswith(".docx"):
        from docx import Document
        return "\n".join(p.text for p in Document(file).paragraphs)
    return file.read().decode("utf-8", errors="ignore")


def prepare_images(files):
    """Shrink phone screenshots so they upload fast and fit Claude's limits."""
    import base64
    from PIL import Image
    out = []
    for f in files:
        img = Image.open(f).convert("RGB")
        img.thumbnail((1568, 1568))
        buf = BytesIO()
        img.save(buf, format="JPEG", quality=85)
        out.append(base64.b64encode(buf.getvalue()).decode())
    return out


def fetch_job_from_link():
    url = st.session_state.get("job_url", "").strip()
    st.session_state["fetch_msg"] = ""
    if not url:
        return
    try:
        from bs4 import BeautifulSoup
        html = requests.get(url, timeout=10, headers={"User-Agent": "Mozilla/5.0"}).text
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style", "nav", "header", "footer"]):
            tag.decompose()
        text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
        if len(text) < 300:
            raise ValueError("too short")
        st.session_state["job_text"] = text[:15000]
        st.session_state["fetch_msg"] = "ok"
    except Exception:
        st.session_state["fetch_msg"] = "fail"


# ================= Editable tables <-> profile =================

ROW_COLS = ["Section", "Title or project", "Company or role", "Location", "Dates", "Bullet"]


def profile_to_rows(profile):
    rows = []
    for job in profile["experience"]:
        for b in job["bullets"] or [{"text": ""}]:
            rows.append(["Experience", job.get("title", ""), job.get("company", ""), job.get("location", ""),
                         job.get("dates", ""), b["text"]])
    for proj in profile["projects"]:
        for b in proj["bullets"] or [{"text": ""}]:
            rows.append(["Project", proj.get("name", ""), proj.get("role", ""), "", proj.get("dates", ""), b["text"]])
    if not rows:
        rows.append(["Experience", "", "", "", "", ""])
    return pd.DataFrame(rows, columns=ROW_COLS)


def rows_to_items(df):
    exp, proj, index = [], [], {}
    for _, r in df.fillna("").iterrows():
        section = r["Section"] or "Experience"
        key = (section, r["Title or project"], r["Company or role"], r["Location"], r["Dates"])
        if not any(str(x).strip() for x in key[1:]) and not str(r["Bullet"]).strip():
            continue
        if key not in index:
            if section == "Project":
                item = {"name": key[1], "role": key[2], "dates": key[4], "bullets": []}
                proj.append(item)
            else:
                item = {"title": key[1], "company": key[2], "location": key[3], "dates": key[4], "bullets": []}
                exp.append(item)
            index[key] = item
        if str(r["Bullet"]).strip():
            index[key]["bullets"].append({"text": str(r["Bullet"]).strip()})
    return exp, proj


def clean_records(df):
    return [{k: str(v).strip() for k, v in r.items()} for r in df.fillna("").to_dict("records")
            if any(str(v).strip() for v in r.values())]


def lines(text):
    return [x.strip() for x in text.splitlines() if x.strip()]


# ================= PDF =================

styles = {
    "name": ParagraphStyle("name", fontName="Times-Bold", fontSize=20, leading=24, alignment=TA_CENTER),
    "contact": ParagraphStyle("contact", fontName="Times-Roman", fontSize=10, leading=13, alignment=TA_CENTER),
    "section": ParagraphStyle("section", fontName="Times-Bold", fontSize=11.5, leading=14, spaceBefore=8),
    "body": ParagraphStyle("body", fontName="Times-Roman", fontSize=10, leading=12.5),
    "left": ParagraphStyle("left", fontName="Times-Bold", fontSize=10.5, leading=13),
    "right": ParagraphStyle("right", fontName="Times-Roman", fontSize=10, leading=13, alignment=TA_RIGHT),
    "sub": ParagraphStyle("sub", fontName="Times-Italic", fontSize=10, leading=12),
}


def link(url):
    label = url.replace("https://", "").replace("http://", "").replace("www.", "").rstrip("/")
    href = url if url.startswith("http") else "https://" + url
    return f'<a href="{escape(href)}"><u>{escape(label)}</u></a>'


def section(story, title):
    story.append(Paragraph(title.upper(), styles["section"]))
    story.append(HRFlowable(width="100%", thickness=0.7, color=colors.black, spaceBefore=1, spaceAfter=4))


def row(left, right):
    t = Table([[Paragraph(left, styles["left"]), Paragraph(right, styles["right"])]],
              colWidths=[5.2 * inch, 2.3 * inch])
    t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                           ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def bullet_list(items):
    return ListFlowable([ListItem(Paragraph(escape(b), styles["body"])) for b in items],
                        bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8)


def pick(prefix, items, ids):
    chosen = []
    for x in ids:
        num = x[len(prefix):]
        if x.startswith(prefix) and num.isdigit() and int(num) < len(items):
            chosen.append(items[int(num)])
    return chosen


def build_pdf(profile, catalog, bullet_ids, skills, result):
    picked = {bid: catalog[bid] for bid in bullet_ids if bid in catalog}
    c = profile["contact"]
    story = [Paragraph(escape(c.get("name", "")), styles["name"])]

    line1 = " | ".join(escape(c[k]) for k in ["phone", "email", "location"] if c.get(k))
    line2 = " | ".join(link(c[k]) for k in ["linkedin", "github", "website"] if c.get(k))
    for line in [line1, line2]:
        if line:
            story.append(Paragraph(line, styles["contact"]))

    summary = profile["summaries"].get(result.get("summary_key"), "") or next(iter(profile["summaries"].values()), "")
    if summary:
        section(story, "Summary")
        story.append(Paragraph(escape(" ".join(summary.split())), styles["body"]))

    if profile["education"]:
        section(story, "Education")
        for e in profile["education"]:
            story.append(row(escape(e.get("degree", "")), escape(str(e.get("graduation", "")))))
            sub = ", ".join(x for x in [e.get("school", ""), e.get("location", "")] if x)
            story.append(Paragraph(escape(sub), styles["sub"]))
            story.append(Spacer(1, 3))

    def has_any(prefix, items):
        return any(f"{prefix}{i}_{j}" in picked for i, it in enumerate(items) for j in range(len(it["bullets"])))

    if has_any("exp", profile["experience"]):
        section(story, "Experience")
        for i, job in enumerate(profile["experience"]):
            bullets = [picked[f"exp{i}_{j}"] for j in range(len(job["bullets"])) if f"exp{i}_{j}" in picked]
            if bullets:
                story.append(row(escape(job.get("title", "")), escape(job.get("dates", ""))))
                sub = ", ".join(x for x in [job.get("company", ""), job.get("location", "")] if x)
                story.append(Paragraph(escape(sub), styles["sub"]))
                story.append(bullet_list(bullets))
                story.append(Spacer(1, 4))

    if has_any("proj", profile["projects"]):
        section(story, "Projects")
        for i, proj in enumerate(profile["projects"]):
            bullets = [picked[f"proj{i}_{j}"] for j in range(len(proj["bullets"])) if f"proj{i}_{j}" in picked]
            if bullets:
                role = f' <font name="Times-Italic">| {escape(proj["role"])}</font>' if proj.get("role") else ""
                story.append(row(f'{escape(proj.get("name", ""))}{role}', escape(proj.get("dates", ""))))
                story.append(bullet_list(bullets))
                story.append(Spacer(1, 4))

    leads = pick("lead", profile["leadership"], result.get("leadership_ids", []))
    if leads:
        section(story, "Leadership")
        story.append(bullet_list([l["text"] for l in leads]))

    if skills:
        section(story, "Skills")
        story.append(Paragraph(escape(", ".join(skills)), styles["body"]))

    certs = pick("cert", profile["certifications"], result.get("cert_ids", []))
    if certs:
        section(story, "Certifications")
        story.append(Paragraph(escape("; ".join(
            f'{x.get("name", "")}, {x.get("issuer", "")} ({x.get("date", "")})' for x in certs)), styles["body"]))

    affs = pick("aff", profile["affiliations"], result.get("affiliation_ids", []))
    if affs:
        section(story, "Affiliations")
        story.append(Paragraph(escape(", ".join(a["name"] for a in affs)), styles["body"]))

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.5 * inch, rightMargin=0.5 * inch,
                            topMargin=0.45 * inch, bottomMargin=0.45 * inch,
                            title=f'{c.get("name", "")} Resume', author=c.get("name", ""))
    doc.build(story)
    return buffer.getvalue(), doc.page


def pdf_images(pdf_bytes):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(pdf_bytes)
    images = [pdf[i].render(scale=1.6).to_pil() for i in range(len(pdf))]
    pdf.close()
    return images


# ================= App state =================

ss = st.session_state
ss.setdefault("step", 1)
ss.setdefault("profile", None)
ss.setdefault("result", None)
ss.setdefault("job_text", "")
ss.setdefault("fetch_msg", "")
ss.setdefault("draft", None)
ss.setdefault("uid", 0)
ss.setdefault("scroll_top", False)


def go(step):
    ss["step"] = step
    ss["scroll_top"] = True


def start_over():
    for k in list(ss.keys()):
        del ss[k]


def new_id():
    ss["uid"] += 1
    return ss["uid"]


def make_draft(p):
    return {
        "exp": [{**e, "_id": new_id()} for e in p["experience"]],
        "proj": [{**x, "_id": new_id()} for x in p["projects"]],
        "edu": [{**e, "_id": new_id()} for e in p["education"]],
    }


# Jump to the top of the page whenever the step changes (helps a lot on phones).
if ss["scroll_top"]:
    js = """<script>
      ['[data-testid="stMain"]', 'section.main', '[data-testid="stAppViewContainer"]']
        .forEach(sel => { const el = document.querySelector(sel); if (el) el.scrollTo(0, 0); });
      window.scrollTo(0, 0);
    </script>"""
    try:
        st.html(js, unsafe_allow_javascript=True)
    except TypeError:
        pass
    ss["scroll_top"] = False

st.title("Resume Tailor")
st.caption("Tailor your resume to any job, using only your real experience.")
st.progress(ss["step"] / 4, text=["You", "Review", "Job", "Preview"][ss["step"] - 1] + f"  ·  Step {ss['step']} of 4")

# ---------- Step 1: You ----------
if ss["step"] == 1:
    st.header("1. Tell us about you")
    name = st.text_input("Your full name")
    method = st.radio("Add your experience", ["Upload my resume", "Upload screenshots"],
                      captions=["PDF or Word file", "Easiest on a phone. Screenshot your LinkedIn or resume."])

    source_text, shots = "", []
    if method == "Upload my resume":
        up = st.file_uploader("PDF or Word file", type=["pdf", "docx", "txt"], label_visibility="collapsed")
        if up is not None:
            try:
                source_text = read_upload(up)
                if len(source_text.strip()) < 50:
                    st.warning("I couldn't read that file. Try screenshots instead.")
            except Exception:
                st.warning("I couldn't read that file. Try screenshots instead.")
    else:
        st.caption("Open each job's 'see more' first, then scroll and screenshot. Up to 8.")
        shots = st.file_uploader("Screenshots", type=["png", "jpg", "jpeg", "webp"],
                                 accept_multiple_files=True, label_visibility="collapsed") or []
        if len(shots) > 8:
            st.warning("Please upload 8 screenshots or fewer.")

    if st.button("Continue", type="primary", width="stretch"):
        has_input = len(source_text.strip()) >= 50 or 0 < len(shots) <= 8
        if not name.strip():
            st.warning("Please add your name.")
        elif not has_input:
            st.warning("Please upload your resume or screenshots first.")
        else:
            try:
                with st.spinner("Reading your experience..."):
                    profile = parse_resume(images=prepare_images(shots)) if shots else parse_resume(source_text)
                profile["contact"]["name"] = name.strip()
                ss["profile"], ss["draft"] = profile, None
                go(2)
                st.rerun()
            except Exception as e:
                st.error("Something went wrong reading your experience. Please try again.")
                st.caption(str(e))

# ---------- Step 2: Review and edit ----------
elif ss["step"] == 2:
    p = ss["profile"]
    if ss["draft"] is None:
        ss["draft"] = make_draft(p)
    d = ss["draft"]

    st.header("2. Review your info")
    st.write("Tap any section to check it. Fix, add, or remove anything.")

    with st.expander("Contact info"):
        c = p["contact"]
        contact = {k: st.text_input(label, c.get(k, ""), key=f"c_{k}") for k, label in [
            ("name", "Name"), ("email", "Email"), ("phone", "Phone"), ("location", "City, State"),
            ("linkedin", "LinkedIn URL"), ("github", "GitHub URL"), ("website", "Website")]}
        summary = st.text_area("Summary", next(iter(p["summaries"].values()), ""), key="c_summary")

    def item_cards(items, kind):
        for it in items:
            uid = it["_id"]
            if kind == "exp":
                label = " at ".join(x for x in [it.get("title", ""), it.get("company", "")] if x) or "New job"
            else:
                label = it.get("name", "") or "New project"
            with st.expander(label):
                if kind == "exp":
                    st.text_input("Job title", it.get("title", ""), key=f"{uid}_title")
                    st.text_input("Company", it.get("company", ""), key=f"{uid}_company")
                    st.text_input("Location", it.get("location", ""), key=f"{uid}_location")
                else:
                    st.text_input("Project name", it.get("name", ""), key=f"{uid}_name")
                    st.text_input("Your role", it.get("role", ""), key=f"{uid}_role")
                st.text_input("Dates", it.get("dates", ""), key=f"{uid}_dates")
                st.text_area("Bullets (one per line)", "\n".join(b["text"] for b in it.get("bullets", [])),
                             key=f"{uid}_bullets", height=160)
                if st.button("Remove", key=f"{uid}_remove"):
                    items.remove(it)
                    st.rerun()

    st.subheader("Experience")
    item_cards(d["exp"], "exp")
    if st.button("+ Add a job", width="stretch"):
        d["exp"].append({"_id": new_id(), "bullets": []})
        st.rerun()

    st.subheader("Projects")
    item_cards(d["proj"], "proj")
    if st.button("+ Add a project", width="stretch"):
        d["proj"].append({"_id": new_id(), "bullets": []})
        st.rerun()

    st.subheader("Education")
    for e in d["edu"]:
        uid = e["_id"]
        with st.expander(e.get("degree", "") or "New education"):
            st.text_input("Degree", e.get("degree", ""), key=f"{uid}_degree")
            st.text_input("School", e.get("school", ""), key=f"{uid}_school")
            st.text_input("Location", e.get("location", ""), key=f"{uid}_elocation")
            st.text_input("Graduation", str(e.get("graduation", "")), key=f"{uid}_grad")
            if st.button("Remove", key=f"{uid}_eremove"):
                d["edu"].remove(e)
                st.rerun()
    if st.button("+ Add education", width="stretch"):
        d["edu"].append({"_id": new_id()})
        st.rerun()

    with st.expander("Skills and extras"):
        skills_text = st.text_area("Skills (separate with commas)", ", ".join(all_skills(p)), key="x_skills")
        cert_text = st.text_area("Certifications (one per line: Name, Issuer, Date)",
                                 "\n".join(", ".join(x for x in [c.get("name", ""), c.get("issuer", ""),
                                                                 c.get("date", "")] if x)
                                           for c in p["certifications"]), key="x_certs")
        lead_text = st.text_area("Leadership (one per line)", "\n".join(l["text"] for l in p["leadership"]),
                                 key="x_lead")
        aff_text = st.text_area("Affiliations (one per line)", "\n".join(a["name"] for a in p["affiliations"]),
                                key="x_aff")

    if st.button("Save and continue", type="primary", width="stretch"):
        g = lambda key: str(ss.get(key, "")).strip()
        exp = [{"title": g(f"{e['_id']}_title"), "company": g(f"{e['_id']}_company"),
                "location": g(f"{e['_id']}_location"), "dates": g(f"{e['_id']}_dates"),
                "bullets": [{"text": t} for t in lines(g(f"{e['_id']}_bullets"))]} for e in d["exp"]]
        proj = [{"name": g(f"{x['_id']}_name"), "role": g(f"{x['_id']}_role"), "dates": g(f"{x['_id']}_dates"),
                 "bullets": [{"text": t} for t in lines(g(f"{x['_id']}_bullets"))]} for x in d["proj"]]
        edu = [{"degree": g(f"{e['_id']}_degree"), "school": g(f"{e['_id']}_school"),
                "location": g(f"{e['_id']}_elocation"), "graduation": g(f"{e['_id']}_grad")} for e in d["edu"]]
        exp = [e for e in exp if e["title"] or e["company"] or e["bullets"]]
        proj = [x for x in proj if x["name"] or x["bullets"]]
        edu = [e for e in edu if any(e.values())]
        if not any(x["bullets"] for x in exp + proj):
            st.warning("Add at least one bullet before continuing.")
        else:
            certs = []
            for line in lines(cert_text):
                parts = [x.strip() for x in line.split(",")]
                certs.append({"name": parts[0], "issuer": parts[1] if len(parts) > 1 else "",
                              "date": ", ".join(parts[2:]) if len(parts) > 2 else ""})
            ss["profile"] = normalize({
                "contact": {k: v.strip() for k, v in contact.items()},
                "summaries": {"default": summary.strip()},
                "experience": exp, "projects": proj, "education": edu,
                "skills": {"all": [x.strip() for x in skills_text.split(",") if x.strip()]},
                "certifications": certs,
                "leadership": [{"text": t} for t in lines(lead_text)],
                "affiliations": [{"name": t} for t in lines(aff_text)],
            })
            ss["draft"] = None
            go(3)
            st.rerun()
    if st.button("Back", width="stretch"):
        go(1)
        st.rerun()

# ---------- Step 3: Job ----------
elif ss["step"] == 3:
    st.header("3. The job")
    st.text_input("Job posting link", key="job_url", placeholder="https://...")
    st.button("Get job from link", on_click=fetch_job_from_link)
    if ss["fetch_msg"] == "ok":
        st.success("Got it. Check the text below and trim anything that isn't the job description.")
    elif ss["fetch_msg"] == "fail":
        st.warning("That site blocked us, which is common for LinkedIn and Indeed. "
                   "Please copy the job description and paste it below.")
    st.text_area("Job description (paste here, or it fills in from the link)", key="job_text", height=280)

    if st.button("Tailor my resume", type="primary", width="stretch"):
        if len(ss["job_text"].strip()) < 100:
            st.warning("Please add the job description first.")
        else:
            try:
                with st.spinner("Picking your best experience for this job..."):
                    _, catalog_text = build_catalog(ss["profile"])
                    result = tailor(ss["profile"], catalog_text, ss["job_text"])
                ss["result"] = result
                ss["sel_bullets"] = result.get("bullet_ids", [])
                ss["sel_skills"] = [s for s in result.get("skills_to_highlight", []) if s in all_skills(ss["profile"])]
                go(4)
                st.rerun()
            except Exception as e:
                st.error("Something went wrong. Please try again.")
                st.caption(str(e))
    if st.button("Back to my info", width="stretch"):
        ss["draft"] = None
        go(2)
        st.rerun()

# ---------- Step 4: Preview ----------
elif ss["step"] == 4:
    p, result = ss["profile"], ss["result"]
    catalog, _ = build_catalog(p)
    st.header("4. Preview")
    st.write(f"**Why these choices:** {result.get('why', '')}")

    with st.expander("Adjust what's included", expanded=False):
        st.caption("Changes update the preview right away. No AI needed.")
        st.multiselect("Bullets", options=list(catalog.keys()), key="sel_bullets",
                       format_func=lambda bid: catalog[bid][:110])
        st.multiselect("Skills", options=all_skills(p), key="sel_skills")

    valid = [b for b in ss["sel_bullets"] if b in catalog]
    pdf_bytes, pages = build_pdf(p, catalog, valid, ss["sel_skills"], result)
    if pages > 1:
        st.warning(f"This is {pages} pages. Most early-career resumes are 1 page. "
                   "Open 'Adjust what's included' and remove a few bullets.")
    try:
        for img in pdf_images(pdf_bytes):
            st.image(img, width="stretch")
    except Exception:
        st.info("Preview isn't available right now, but your download will work.")

    name = (p["contact"].get("name") or "resume").replace(" ", "_")
    st.download_button("Download PDF", pdf_bytes, file_name=f"{name}_tailored.pdf",
                       mime="application/pdf", type="primary", width="stretch")

    b1 = b2 = b3 = st
    if b1.button("Tailor for another job", width="stretch"):
        ss["job_text"], ss["job_url"], ss["fetch_msg"] = "", "", ""
        go(3)
        st.rerun()
    if b2.button("Edit my info", width="stretch"):
        ss["draft"] = None
        go(2)
        st.rerun()
    if b3.button("Delete everything", width="stretch"):
        start_over()
        st.rerun()
    st.caption("Nothing is saved on our side. Download your PDF to keep it. Saving online is coming in V2.")

# ================= Feedback =================
st.divider()
with st.expander("💬 Share feedback with Tessy"):
    with st.form("feedback", clear_on_submit=True):
        rating = st.radio("How useful was this?", ["Very useful", "Somewhat useful", "Not useful"], horizontal=True)
        confused = st.text_area("What confused you?")
        wish = st.text_area("What would you want next?")
        who = st.text_input("Your name (optional)")
        if st.form_submit_button("Send feedback", width="stretch"):
            payload = {"rating": rating, "confused": confused, "wish": wish, "name": who, "step": ss.get("step")}
            url = st.secrets.get("FEEDBACK_URL", "")
            try:
                if not url:
                    raise ValueError("no feedback url")
                r = requests.post(url, json=payload, headers={"Accept": "application/json"}, timeout=10)
                r.raise_for_status()
                st.success("Thank you! Your feedback was sent to Tessy.")
            except Exception:
                st.warning("Couldn't send it automatically. Please copy this and send it to Tessy:")
                st.code(json.dumps(payload, indent=2))
